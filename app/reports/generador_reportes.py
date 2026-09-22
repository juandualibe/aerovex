"""Módulo de generación de reportes técnicos para AEROVEX.

Permite exportar los resultados de misiones a hojas de cálculo Excel (.xlsx)
y documentos ejecutivos en PDF con métricas de carga ganadera.
"""

from datetime import datetime, timezone

from openpyxl import Workbook  # type: ignore[import-untyped]
from openpyxl.styles import (  # type: ignore[import-untyped]
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from reportlab.lib import colors  # type: ignore[import-untyped]
from reportlab.lib.pagesizes import letter  # type: ignore[import-untyped]
from reportlab.lib.styles import (  # type: ignore[import-untyped]
    ParagraphStyle,
    getSampleStyleSheet,
)
from reportlab.lib.units import inch  # type: ignore[import-untyped]
from reportlab.platypus import (  # type: ignore[import-untyped]
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


class GeneradorReportes:
    """Genera reportes en formato Excel y PDF a partir de una misión completada."""

    @staticmethod
    def exportar_excel(datos_mision: dict, ruta_salida: str) -> str:
        """Crea un archivo Excel estructurado con el resumen y detalle por foto."""
        wb = Workbook()
        ws = wb.active
        ws.title = "Resumen de Censo"

        fuente_titulo = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
        fuente_encabezado = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        fuente_datos = Font(name="Calibri", size=10)

        fill_verde_oscuro = PatternFill(
            start_color="1E3A8A", end_color="1E3A8A", fill_type="solid"
        )
        fill_azul = PatternFill(
            start_color="2563EB", end_color="2563EB", fill_type="solid"
        )
        fill_zebra = PatternFill(
            start_color="F1F5F9", end_color="F1F5F9", fill_type="solid"
        )

        borde_fino = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1"),
        )

        # Encabezado principal
        ws.merge_cells("A1:D1")
        ws["A1"] = "AEROVEX - REPORTE DE CONTEO GANADERO"
        ws["A1"].font = fuente_titulo
        ws["A1"].fill = fill_verde_oscuro
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 35

        # Datos generales
        potrero = datos_mision.get("potrero_nombre", "General")
        superficie = float(datos_mision.get("superficie_ha", 0.0))
        total = int(datos_mision.get("total_cabezas", 0))
        densidad = round(total / superficie, 2) if superficie > 0 else 0.0

        fecha_defecto = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M")
        fecha_mision = datos_mision.get("fecha", fecha_defecto)

        metadatos = [
            ("Fecha de Misión:", fecha_mision),
            ("Potrero Censado:", f"{potrero} ({superficie} Ha)"),
            ("Total Cabezas Detectadas:", total),
            ("Carga Animal Calculada:", f"{densidad} cabezas/Ha"),
        ]

        fila_actual = 3
        for etiqueta, valor in metadatos:
            ws.cell(row=fila_actual, column=1, value=etiqueta).font = Font(
                name="Calibri", bold=True
            )
            ws.cell(row=fila_actual, column=2, value=valor).font = fuente_datos
            fila_actual += 1

        fila_actual += 1

        # Tabla de detalle por archivo
        columnas = ["N°", "Archivo de Vuelo", "Animales Censados", "Estado"]
        for col_idx, col_nombre in enumerate(columnas, start=1):
            celda = ws.cell(row=fila_actual, column=col_idx, value=col_nombre)
            celda.font = fuente_encabezado
            celda.fill = fill_azul
            celda.alignment = Alignment(horizontal="center", vertical="center")

        ws.row_dimensions[fila_actual].height = 24
        fila_actual += 1

        fotos = datos_mision.get("fotos", [])
        for idx, f in enumerate(fotos, start=1):
            c1 = ws.cell(row=fila_actual, column=1, value=idx)
            c2 = ws.cell(row=fila_actual, column=2, value=f.get("archivo", ""))
            c3 = ws.cell(row=fila_actual, column=3, value=f.get("cantidad", 0))
            c4 = ws.cell(row=fila_actual, column=4, value="Procesado OK")

            for celda in (c1, c2, c3, c4):
                celda.font = fuente_datos
                celda.border = borde_fino
                if idx % 2 == 0:
                    celda.fill = fill_zebra

            c1.alignment = Alignment(horizontal="center")
            c3.alignment = Alignment(horizontal="center")
            c4.alignment = Alignment(horizontal="center")
            fila_actual += 1

        # Ancho de columnas
        ws.column_dimensions["A"].width = 8
        ws.column_dimensions["B"].width = 38
        ws.column_dimensions["C"].width = 20
        ws.column_dimensions["D"].width = 18

        wb.save(ruta_salida)
        return ruta_salida

    @staticmethod
    def exportar_pdf(datos_mision: dict, ruta_salida: str) -> str:
        """Genera un reporte PDF con diseño ejecutivo."""
        doc = SimpleDocTemplate(
            ruta_salida,
            pagesize=letter,
            rightMargin=40,
            leftMargin=40,
            topMargin=40,
            bottomMargin=40,
        )

        estilos = getSampleStyleSheet()
        estilo_titulo = ParagraphStyle(
            "TituloDoc",
            parent=estilos["Heading1"],
            fontSize=20,
            leading=24,
            textColor=colors.HexColor("#0F172A"),
            spaceAfter=6,
        )
        estilo_subtitulo = ParagraphStyle(
            "SubtituloDoc",
            parent=estilos["Normal"],
            fontSize=10,
            leading=13,
            textColor=colors.HexColor("#64748B"),
            spaceAfter=15,
        )
        estilo_texto = ParagraphStyle(
            "Cuerpo",
            parent=estilos["Normal"],
            fontSize=10,
            leading=14,
            textColor=colors.HexColor("#1E293B"),
        )

        elementos = []

        elementos.append(
            Paragraph("AEROVEX | Reporte de Relevamiento Ganadero", estilo_titulo)
        )
        elementos.append(
            Paragraph(
                "Sistema Autónomo de Conteo y Monitoreo con Vehículos Aéreos No Tripulados",
                estilo_subtitulo,
            )
        )
        elementos.append(Spacer(1, 10))

        potrero = datos_mision.get("potrero_nombre", "General")
        superficie = float(datos_mision.get("superficie_ha", 0.0))
        total = int(datos_mision.get("total_cabezas", 0))
        densidad = round(total / superficie, 2) if superficie > 0 else 0.0

        fecha_defecto = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M")
        fecha_mision = datos_mision.get("fecha", fecha_defecto)

        tabla_metricas_datos = [
            ["Fecha de Inspección", fecha_mision],
            ["Potrero Evaluado", f"{potrero} ({superficie} Hectáreas)"],
            ["Total Cabezas Censadas", str(total)],
            ["Carga Animal", f"{densidad} animales/Ha"],
        ]

        t_resumen = Table(tabla_metricas_datos, colWidths=[2.2 * inch, 4.5 * inch])
        t_resumen.setStyle(
            TableStyle([
                ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#F8FAFC")),
                ("TEXTCOLOR", (0, 0), (-1, -1), colors.HexColor("#0F172A")),
                ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
            ])
        )
        elementos.append(t_resumen)
        elementos.append(Spacer(1, 20))

        elementos.append(
            Paragraph("<b>Desglose por Fotografía de Captura</b>", estilo_texto)
        )
        elementos.append(Spacer(1, 8))

        fotos = datos_mision.get("fotos", [])
        filas_detalle = [["N°", "Nombre de Captura", "Detecciones", "Estado"]]
        for idx, f in enumerate(fotos, start=1):
            filas_detalle.append([
                str(idx),
                str(f.get("archivo", "")),
                str(f.get("cantidad", 0)),
                "OK",
            ])

        t_detalle = Table(
            filas_detalle,
            colWidths=[0.5 * inch, 3.8 * inch, 1.3 * inch, 1.1 * inch],
        )
        t_detalle.setStyle(
            TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E293B")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("ALIGN", (0, 0), (0, -1), "CENTER"),
                ("ALIGN", (2, 0), (3, -1), "CENTER"),
                (
                    "ROWBACKGROUNDS",
                    (0, 1),
                    (-1, -1),
                    [colors.white, colors.HexColor("#F1F5F9")],
                ),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
            ])
        )
        elementos.append(t_detalle)

        doc.build(elementos)
        return ruta_salida